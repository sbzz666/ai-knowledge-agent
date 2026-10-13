from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import fitz
from docx import Document


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedPage:
    page: int | None
    text: str


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


def parse_document(filename: str, data: bytes, *, max_pages: int = 80) -> list[ParsedPage]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentParseError("仅支持 PDF、DOCX、TXT 和 Markdown 文件")
    if not data:
        raise DocumentParseError("文件内容为空")

    if suffix == ".pdf":
        if not data.startswith(b"%PDF"):
            raise DocumentParseError("文件扩展名为 PDF，但内容不是有效 PDF")
        try:
            document = fitz.open(stream=data, filetype="pdf")
        except Exception as exc:
            raise DocumentParseError("PDF 无法解析") from exc
        if document.page_count > max_pages:
            raise DocumentParseError(f"PDF 不能超过 {max_pages} 页")
        pages = [ParsedPage(index + 1, page.get_text("text")) for index, page in enumerate(document)]
    elif suffix == ".docx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise DocumentParseError("文件扩展名为 DOCX，但内容不是有效 Word 文档")
        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError("DOCX 无法解析") from exc
        blocks = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        pages = [ParsedPage(None, "\n".join(blocks))]
    else:
        try:
            pages = [ParsedPage(None, data.decode("utf-8-sig"))]
        except UnicodeDecodeError as exc:
            raise DocumentParseError("文本文件必须使用 UTF-8 编码") from exc

    cleaned = [ParsedPage(page.page, _clean(page.text)) for page in pages if _clean(page.text)]
    if not cleaned:
        raise DocumentParseError("没有提取到可用文字；扫描版 PDF 暂不支持 OCR")
    return cleaned


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

